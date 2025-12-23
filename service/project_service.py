from openpyxl import load_workbook

from db.project_dao import ProjectDao
from service.queue_service import QueueService


class ProjectService:
    def __init__(self):
        self.db = ProjectDao()
        self.queue_service = QueueService()

    def insert_project(self, project_file):
        wb = load_workbook(project_file)
        sheet = wb["Sheet1"]
        for row in sheet.iter_rows(values_only=True, min_row=2):
            if row[4].strip() != "" or row[5].strip() != "":
                dict_project = {
                    "customer_name": row[0],
                    "project_name": row[1],
                    "project_description": row[2],
                    "reference": row[3],
                    "tech_stack": row[4],
                    "service": row[5].replace("&amp;", "&"),
                }
                project_id = self.db.insert_project(dict_project)
                self.queue_service.publish_event("project.load", {"project_id": project_id})
        self.queue_service.close_connection()
